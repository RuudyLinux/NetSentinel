"""One-off generator for NetSentinel AI's project documentation (.docx).

Run: uv run --with python-docx python Documentation/generate_docx.py
Source of truth: README.md, docs/architecture.md, docs/demo.md, SECURITY.md,
rules/cis/cisco-ios-v8.yaml, and the live application (RBAC matrix, page list,
API surface) as verified during QA. Not generated from the aspirational
latest.md spec — that covers SP2-SP6 work not yet built.
"""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Pt, RGBColor, Cm

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "Documentation" / "NetSentinel AI - Project Documentation.docx"

NAVY = RGBColor(0x0F, 0x17, 0x2A)
SLATE = RGBColor(0x33, 0x41, 0x55)
SKY = RGBColor(0x03, 0x69, 0xA1)


def set_cell_shading(cell, hex_color: str) -> None:
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    cell._tc.get_or_add_tcPr().append(shd)


def style_table(table) -> None:
    table.style = "Light Grid Accent 1"
    for cell in table.rows[0].cells:
        set_cell_shading(cell, "0F172A")
        for p in cell.paragraphs:
            for r in p.runs:
                r.font.bold = True
                r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)


def h1(doc, text):
    p = doc.add_heading(text, level=1)
    for r in p.runs:
        r.font.color.rgb = NAVY
    return p


def h2(doc, text):
    p = doc.add_heading(text, level=2)
    for r in p.runs:
        r.font.color.rgb = SLATE
    return p


def body(doc, text, bold=False, italic=False):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold = bold
    r.italic = italic
    return p


def bullets(doc, items):
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def mono_block(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(text)
    r.font.name = "Consolas"
    r.font.size = Pt(9)
    set_cell_shading  # noqa (silence unused if refactored later)
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), "F1F5F9")
    pPr.append(shd)
    return p


doc = Document()

# ---- base styles ----
normal = doc.styles["Normal"]
normal.font.name = "Calibri"
normal.font.size = Pt(11)
normal.font.color.rgb = RGBColor(0x1A, 0x1A, 0x1A)

for sec in doc.sections:
    sec.left_margin = Cm(2.2)
    sec.right_margin = Cm(2.2)

# ---- Title page ----
title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.paragraph_format.space_before = Pt(160)
r = title.add_run("NetSentinel AI")
r.font.size = Pt(40)
r.font.bold = True
r.font.color.rgb = NAVY

subtitle = doc.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = subtitle.add_run("Project Documentation — Slice 1 (the spine)")
r.font.size = Pt(16)
r.font.color.rgb = SKY

tagline = doc.add_paragraph()
tagline.alignment = WD_ALIGN_PARAGRAPH.CENTER
tagline.paragraph_format.space_before = Pt(24)
r = tagline.add_run(
    "A network security compliance auditor: upload a device configuration and get an\n"
    "explainable posture score, evidenced findings, remediation guidance, and a PDF\n"
    "report — with every step traceable back to specific configuration lines."
)
r.font.size = Pt(11)
r.font.italic = True
r.font.color.rgb = SLATE

doc.add_page_break()

# ---- Table of contents (Word field — press F9 / "Update Field" after opening) ----
h1(doc, "Table of Contents")
body(doc, "This field populates when the document is opened in Microsoft Word "
          "(right-click → Update Field, or press F9).")
p = doc.add_paragraph()
run = p.add_run()
fld = OxmlElement("w:fldSimple")
fld.set(qn("w:instr"), 'TOC \\o "1-2" \\h \\z \\u')
run._r.addnext(fld)
doc.add_page_break()

# ---- 1. Executive summary ----
h1(doc, "1. Executive Summary")
body(doc,
     "NetSentinel AI is a network security compliance auditor. An operator uploads a "
     "device configuration file; the system identifies the vendor and OS with a stated "
     "confidence, parses it into a structured tree, evaluates it against a versioned "
     "compliance rule pack, and produces an explainable, severity-weighted posture "
     "score, evidenced findings, device-specific remediation guidance, and a PDF report "
     "— every step traceable back to the exact configuration line it came from.")
body(doc,
     "This repository currently implements slice 1 (\"the spine\"): one vendor (Cisco "
     "IOS / IOS-XE), one framework (CIS 8.0, sixteen controls) — but the complete "
     "pipeline end to end, with real authentication, RBAC, and provenance, not a demo "
     "shortcut.")
h2(doc, "In scope (slice 1)")
bullets(doc, [
    "Email/password authentication — Argon2id hashing, JWT access + rotating refresh tokens",
    "Six RBAC roles enforced server-side on every endpoint",
    "Configuration upload with validation, secret redaction, and SHA-256 content hashing",
    "Cisco IOS / IOS-XE vendor detection with a confidence score and stated reasons",
    "Cisco IOS block parser producing a typed configuration tree with line numbers",
    "Normalizer mapping parsed facts to sixteen Universal Security Control keys",
    "Deterministic rule engine loading versioned YAML rule packs (CIS, all 16 controls)",
    "Findings with evidence, provenance, and severity; device-specific remediation guidance",
    "Explainable, severity-weighted posture score with coverage",
    "PDF device report",
    "Advisory AI interpretation of configuration the parser doesn't recognize (never changes a score; a human always approves it)",
    "React UI: login, dashboard, devices, discovery, upload, audits, findings, compliance, reports, admin",
])
h2(doc, "Deferred to later sub-projects (SP2-SP6)")
bullets(doc, [
    "Additional vendors and additional compliance frameworks",
    "Adaptive AI training loop (learned mappings feeding back into re-parsing)",
    "Configuration drift detection and what-if hardening",
    "Risk graph visualization",
    "Tamper-evident evidence ledger",
    "SSO / SAML / LDAP",
    "Postgres / Redis / MinIO / Celery containerization",
])
doc.add_page_break()

# ---- 2. Problem statement ----
h1(doc, "2. Problem Statement")
body(doc,
     "Manual network configuration compliance review is slow, inconsistent between "
     "reviewers, and hard to evidence after the fact. Auditors need proof a control was "
     "actually assessed against actual configuration text — not a pass/fail badge with "
     "no trail back to source. Security teams need remediation guidance scoped to the "
     "specific device and vendor, not a generic checklist. NetSentinel AI's spine exists "
     "to make every one of those steps deterministic, evidenced, and repeatable for one "
     "vendor and one framework, as the foundation later phases extend rather than "
     "replace.")
doc.add_page_break()

# ---- 3. Architecture ----
h1(doc, "3. Architecture")
h2(doc, "3.1 Pipeline")
mono_block(doc,
    "POST /configurations/upload\n"
    "        |\n"
    "        v\n"
    "   Sanitizer   -> validate ext/size . detect secrets . redact . SHA-256 . store blob\n"
    "        |\n"
    "        v   RedactedConfig(text, sha256, blob_key, secret_hits)\n"
    "        |\n"
    "POST /audits\n"
    "        |\n"
    "        v\n"
    "   Detector    -> DeviceIdentity(vendor, os, version, confidence, reasons)\n"
    "        |\n"
    "        v\n"
    "   Parser      -> ConfigTree(nodes with line numbers) + ParseWarning[]\n"
    "        |\n"
    "        v\n"
    "   Normalizer  -> ControlSet{key: ControlValue} + UnknownConstruct[]\n"
    "        |\n"
    "        v\n"
    "   RuleEngine  -> ComplianceResult[] (rules loaded from versioned YAML packs)\n"
    "        |\n"
    "        v\n"
    "   FindingBuilder -> Finding[] (FAIL/WARNING only, + remediation lookup)\n"
    "        |\n"
    "        v\n"
    "   Scorer      -> PostureScore(score, coverage, severity counts)\n"
    "        |\n"
    "        v\n"
    "   ReportBuilder -> PDF (on demand, POST /audits/{id}/report)")
body(doc,
     "Every stage is a pure function from the previous stage's typed output to its own. "
     "No stage opens a database session, reads a file, or makes an HTTP call. "
     "Persistence and orchestration live exclusively in "
     "backend/app/services/audit/runner.py (run_audit(...) -> AuditRun), which calls the "
     "stages in order and writes results in one transaction.")
body(doc,
     "This is what keeps vendor and framework additions cheap: the parser knows Cisco "
     "syntax and nothing about CIS; the rule engine knows CIS and nothing about Cisco. "
     "Adding a vendor touches app/services/detection/ and app/services/parsing/ only. "
     "Adding a framework touches rules/ only — which is data, not code.")

h2(doc, "3.2 Where the AI interpreter attaches")
body(doc,
     "When the normalizer meets syntax it cannot map to a control, it records an "
     "UnknownConstruct and continues rather than failing the audit. Today these are "
     "counted and surfaced on the audit detail page; a user with the MAPPING_SUGGEST "
     "permission can ask the configured AI model (openai/gpt-4o-mini via OpenRouter) "
     "what control the line likely sets. The suggestion carries a confidence score and "
     "the exact configuration text as evidence, and a user with MAPPING_APPROVE must "
     "approve or reject it before it is recorded — it never changes the audit's score "
     "and nothing is applied automatically. The adaptive training loop that would feed "
     "approved mappings back into future parsing is SP4 scope, not yet built.")

h2(doc, "3.3 The three environment seams")
body(doc,
     "The development environment has no Docker and the parent spec assumes Postgres, "
     "Redis, MinIO, and Celery in containers. Slice 1 binds three named seams to local "
     "implementations instead, so adopting the containerized stack later is a "
     "configuration change plus one new class each — never a rewrite of calling code.")
table = doc.add_table(rows=1, cols=4)
style_table(table)
hdr = table.rows[0].cells
hdr[0].text, hdr[1].text, hdr[2].text, hdr[3].text = (
    "Seam", "Slice 1 binding", "Target (SP6) binding", "Cost of switch")
seams = [
    ("DATABASE_URL", "sqlite:///./var/netsentinel.db", "postgresql+psycopg://...", "Env var + Alembic run"),
    ("StorageBackend", "LocalFileStorage (./var/blobs/)", "S3Storage (MinIO/S3)", "One class, ~60 lines"),
    ("TaskRunner", "InlineTaskRunner (BackgroundTasks)", "CeleryTaskRunner", "One class, ~40 lines"),
]
for row in seams:
    cells = table.add_row().cells
    for i, val in enumerate(row):
        cells[i].text = val
body(doc,
     "docker-compose.yml at the repo root defines the SP6 target stack today and is "
     "kept accurate even while unused in slice 1 development.")

h2(doc, "3.4 Boundary discipline")
bullets(doc, [
    "No SQLite-specific SQL anywhere in application code",
    "No raw file paths outside LocalFileStorage",
    "No BackgroundTasks import outside InlineTaskRunner",
    "Every model uses SQLAlchemy types that map cleanly to both SQLite and PostgreSQL",
])
body(doc,
     "These rules are what make the seam table above true rather than aspirational — "
     "violate any of them and the \"one class, ~N lines\" switch cost stops holding.")

h2(doc, "3.5 Universal Security Control model")
body(doc,
     "Every normalized fact carries provenance, not just a value — each finding must "
     "trace back to specific configuration lines:")
mono_block(doc,
    "@dataclass(frozen=True)\n"
    "class ControlValue:\n"
    "    value: bool | int | str | list[str] | None\n"
    "    source_lines: list[int]      # 1-indexed lines in the redacted config\n"
    "    excerpt: str                 # redacted configuration text that produced this value\n"
    "    parser_confidence: float     # 1.0 for deterministic parsing\n"
    "    origin: Literal[\"deterministic\", \"inferred\", \"learned\"]")
body(doc,
     "origin is \"deterministic\" throughout slice 1; \"inferred\"/\"learned\" exist so "
     "SP4 can populate them without a schema migration. A control key absent from the "
     "ControlSet means \"the parser did not observe this\"; present with value=None "
     "means \"observed and explicitly unset\" — the rule engine treats both as "
     "NOT_ASSESSABLE, never as a pass.")
doc.add_page_break()

# ---- 4. Authentication & RBAC ----
h1(doc, "4. Authentication & Access Control")
h2(doc, "4.1 Authentication")
bullets(doc, [
    "Email + password login; passwords hashed with Argon2id",
    "JWT access tokens (short-lived) plus rotating refresh tokens",
    "Refresh-token reuse triggers family-wide revocation and an audit event",
    "Forgot-password flow returns an identical response whether or not the account exists (no user enumeration)",
    "No SSO in slice 1 — SSO/SAML/LDAP is SP6 scope",
])
h2(doc, "4.2 Six roles, enforced server-side")
body(doc,
     "Every one of the twelve permissions below is checked on the server for every "
     "request that needs it — the UI hiding a button is presentation only, never the "
     "actual control. This was independently verified: every denied action attempted "
     "during QA (see Section 8) returned HTTP 403 from the API regardless of what the "
     "UI showed.")
table = doc.add_table(rows=1, cols=7)
style_table(table)
hdr = table.rows[0].cells
headers = ["Permission", "Auditor", "CISO", "Network\nEngineer", "Platform\nAdmin", "Security\nAdmin", "Security\nAnalyst"]
for i, h in enumerate(headers):
    hdr[i].text = h
rbac_rows = [
    ("DEVICE_READ", "Y", "Y", "Y", "Y", "Y", "Y"),
    ("DEVICE_WRITE", "-", "-", "-", "Y", "Y", "-"),
    ("CONFIG_UPLOAD", "-", "-", "Y", "Y", "Y", "-"),
    ("AUDIT_RUN", "-", "-", "Y", "Y", "Y", "Y"),
    ("AUDIT_READ", "Y", "Y", "Y", "Y", "Y", "Y"),
    ("FINDING_READ", "Y", "Y", "Y", "Y", "Y", "Y"),
    ("FINDING_TRIAGE", "-", "-", "Y", "Y", "Y", "Y"),
    ("REPORT_GENERATE", "Y", "Y", "Y", "Y", "Y", "Y"),
    ("REPORT_READ", "Y", "Y", "Y", "Y", "Y", "Y"),
    ("USER_ADMIN", "-", "-", "-", "Y", "-", "-"),
    ("MAPPING_SUGGEST", "-", "-", "Y", "Y", "Y", "Y"),
    ("MAPPING_APPROVE", "-", "-", "-", "Y", "Y", "-"),
]
for row in rbac_rows:
    cells = table.add_row().cells
    for i, val in enumerate(row):
        cells[i].text = val
body(doc, "Y = granted, - = not granted. Read out live at Admin > Roles & Permissions.")
doc.add_page_break()

# ---- 5. Compliance rule pack ----
h1(doc, "5. Compliance Rule Pack — CIS 8.0 (Cisco IOS / IOS-XE)")
body(doc,
     "Rule packs are versioned YAML data, not code — adding a rule never touches "
     "Python. The application refuses to start if a rule pack fails schema validation "
     "or references a remediation id that doesn't exist: rule packs fail closed. "
     "Source: rules/cis/cisco-ios-v8.yaml; matching remediation CLI in mappings/cisco/.")
table = doc.add_table(rows=1, cols=4)
style_table(table)
hdr = table.rows[0].cells
hdr[0].text, hdr[1].text, hdr[2].text, hdr[3].text = ("Rule ID", "Control", "Parameter key", "Severity")
rules_data = [
    ("CIS-SSH-001", "SSH management access enabled", "management.ssh.enabled", "HIGH"),
    ("CIS-SSH-002", "SSH protocol version 2 enforced", "management.ssh.version", "HIGH"),
    ("CIS-TELNET-001", "Telnet management disabled", "management.telnet.enabled", "CRITICAL"),
    ("CIS-HTTP-001", "HTTP management server disabled", "management.http.enabled", "HIGH"),
    ("CIS-HTTPS-001", "HTTPS management server used where web management is required", "management.https.enabled", "MEDIUM"),
    ("CIS-SESSION-001", "Idle management sessions time out within 10 minutes", "management.session_timeout", "MEDIUM"),
    ("CIS-MGMTACL-001", "Management access restricted by access list", "management.acl.present", "HIGH"),
    ("CIS-PASSLEN-001", "Minimum password length of at least 8 characters", "auth.password.min_length", "MEDIUM"),
    ("CIS-AAA-001", "AAA authentication enabled", "auth.aaa.enabled", "HIGH"),
    ("CIS-ENABLE-001", "Privileged password stored with a strong hash", "auth.enable_secret.encrypted", "CRITICAL"),
    ("CIS-DEFACCT-001", "No vendor default accounts present", "auth.default_accounts.present", "CRITICAL"),
    ("CIS-LOGLOCAL-001", "Local logging buffer enabled", "logging.local.enabled", "LOW"),
    ("CIS-LOGREMOTE-001", "Remote syslog destination configured", "logging.remote.enabled", "HIGH"),
    ("CIS-NTP-001", "NTP time synchronization configured", "time.ntp.enabled", "MEDIUM"),
    ("CIS-SNMP-001", "Only SNMPv3 in use", "snmp.v3.only", "HIGH"),
    ("CIS-BANNER-001", "Login banner present", "banner.login.present", "LOW"),
]
for row in rules_data:
    cells = table.add_row().cells
    for i, val in enumerate(row):
        cells[i].text = val
body(doc,
     "A control absent from a device's parsed configuration is scored NOT_ASSESSABLE, "
     "not a silent pass — e.g. if SSH is off entirely, CIS-SSH-002 (its version) cannot "
     "be assessed and is reported as such rather than ignored.")
doc.add_page_break()

# ---- 6. Application walkthrough ----
h1(doc, "6. Application Pages")
h2(doc, "6.1 Public")
bullets(doc, ["Sign in", "Forgot password (generic response, no user enumeration)", "404 page"])
h2(doc, "6.2 Authenticated")
bullets(doc, [
    "Dashboard — posture score, device/finding/audit counts, framework scores, risk trend, recent findings",
    "Devices — list and per-device detail (posture score, recent audits)",
    "Discovery — auto-detects the operator's own subnet, live SSH-banner scan (no credentials tried), connect-and-audit form",
    "Configurations — drag-and-drop upload; detects vendor/OS, redacts secrets, runs the audit, redirects to the result",
    "Audits — list and per-audit detail: detection reasons, provenance (rule pack hash, config hash, engine version), unrecognized-construct panel, findings list with the WHY / EVIDENCE / IMPACT / REMEDIATION / VERIFY panel, Generate PDF report",
    "Findings — cross-fleet list with severity and status filters",
    "Compliance — framework list and per-framework detail (pass/fail/warning/not-assessable counts, every rule, every finding)",
    "AI Insights — status page describing what's connected (unrecognized-construct interpretation) and what's still SP4 scope",
    "Reports — history of generated PDFs with download",
])
h2(doc, "6.3 Administration (USER_ADMIN only)")
bullets(doc, [
    "Users — accounts in the organization",
    "Roles & Permissions — the exact matrix in Section 4.2, read live from the server",
    "Audit Logs — append-only security event log (login/logout, uploads, audits run, triage, report generation) with pagination",
    "API Diagnostics — live-checks every endpoint the deployment exposes, including writes, safely round-tripped or probed for their real failure path; a separate opt-in check exercises the paid AI interpretation call",
])
doc.add_page_break()

# ---- 7. Security ----
h1(doc, "7. Security")
h2(doc, "7.1 Redaction guarantee")
body(doc,
     "Every uploaded configuration is redacted on ingest, before persistence. The "
     "sanitizer matches Cisco type-7/cleartext passwords, enable secret hashes, "
     "username secret/password lines, SNMP community strings, key-string/pre-shared-"
     "key/crypto-key blocks, and PEM private-key blocks, replacing each with a "
     "<REDACTED:type> marker and counting it in Configuration.secret_hits.")
body(doc,
     "The redacted text is the only copy that reaches the database, API responses, "
     "evidence excerpts, PDF reports, logs, or the local blob store — nothing "
     "downstream of redaction ever sees the raw bytes.", bold=False)
h2(doc, "7.2 Threat model coverage (slice 1)")
bullets(doc, [
    "T1 — untrusted upload content: extension/size/encoding validation; the whole parsing/normalization pipeline treats configuration text as untrusted throughout",
    "T2 — secret exposure in outputs: redaction on ingest (7.1)",
    "T4 — unauthorized access to findings/reports: server-side RBAC on every endpoint, organization-scoped queries",
    "T6 — audit trail tampering/gaps: append-only AuditEvent log; refresh-token reuse triggers family-wide revocation and an audit event",
])
body(doc, "Deferred to SP4: T3 and T5 — dependent on the AI interpreter and adaptive training loop.")
h2(doc, "7.3 Development-only defaults — replace before real deployment")
bullets(doc, [
    "The seeded demo accounts' shared password, demo-password-1",
    "Settings.jwt_secret's default value — override via the NETSENTINEL_JWT_SECRET environment variable",
])
body(doc, "Neither default is fit for anything beyond a local demo. Slice 1 has no SSO.")
doc.add_page_break()

# ---- 8. QA & verification ----
h1(doc, "8. QA & Verification")
body(doc,
     "A full autonomous QA pass exercised every discovered page and interactive "
     "element across all six roles: forms (valid/invalid/empty), file upload "
     "(valid/invalid type), the live discovery scan, filters, pagination, PDF "
     "generation and download, the AI interpretation round-trip, and direct-URL "
     "access to permission-gated routes for every role. Backend RBAC held in every "
     "case — no authorization bypass was found; every denied action returned HTTP 403 "
     "from the server regardless of what the client showed.")
body(doc, "Issues found during that pass, and their resolution:", bold=True)
table = doc.add_table(rows=1, cols=3)
style_table(table)
hdr = table.rows[0].cells
hdr[0].text, hdr[1].text, hdr[2].text = ("Issue", "Root cause", "Fix")
qa_rows = [
    ("Login succeeded but never redirected to the dashboard",
     "React Query's enabled flag read localStorage directly instead of reactive state, so the current-user fetch never (re)fired after login",
     "Track token presence in React state; flip it on login/logout"),
    ("Logout cleared tokens but left the authenticated UI on screen",
     "Same root cause as above",
     "Same fix"),
    ("404/403 pages appeared stuck loading for ~7 seconds",
     "Correct error UI already existed; the default query retry policy retried unretryable 4xx errors before showing it",
     "Stop retrying 4xx client errors globally"),
    ("\"Mark Remediated\" was shown to roles without the triage permission and failed silently",
     "No client-side permission gate on that action",
     "Gated behind the same RequirePermission pattern used elsewhere"),
    ("PDF report showed a different coverage percentage than the UI for the same audit",
     "Python's default rounding (round-half-to-even) disagreed with the frontend's Math.round (round-half-up) at exact .5 ties",
     "Round-half-up helper used consistently in the PDF report"),
    ("Configurations page fully rendered the upload form for roles without upload permission",
     "No page-level permission gate; only a submit-time server error",
     "Page-level check with a clear \"you don't have permission\" state"),
]
for row in qa_rows:
    cells = table.add_row().cells
    for i, val in enumerate(row):
        cells[i].text = val
body(doc,
     "All fixes verified live against every affected role, with no regression for "
     "roles that do have the relevant permission. Full backend test suite: 265/265 "
     "passing. Frontend: clean TypeScript build, clean lint.")
doc.add_page_break()

# ---- 9. Running it locally ----
h1(doc, "9. Running It Locally")
h2(doc, "Prerequisites")
bullets(doc, ["uv (Python 3.13, pinned automatically — no system Python changes needed)", "Node.js 22+"])
h2(doc, "Backend (from backend/)")
mono_block(doc, "uv sync --dev\nuv run alembic upgrade head\nuv run python -m scripts.seed\nuv run uvicorn app.main:app --reload")
body(doc, "The API serves on http://127.0.0.1:8000; SQLite lives at backend/var/netsentinel.db and uploaded blobs at backend/var/blobs/.")
h2(doc, "Frontend (from frontend/, in a second terminal)")
mono_block(doc, "npm install\nnpm run dev")
body(doc, "Open the printed URL (http://localhost:5173 by default) — its dev server proxies /api to the backend.")
h2(doc, "Demo accounts")
body(doc, "Seeding creates six users, all sharing the password demo-password-1:")
table = doc.add_table(rows=1, cols=2)
style_table(table)
hdr = table.rows[0].cells
hdr[0].text, hdr[1].text = ("Email", "Role")
for email, role in [
    ("admin@netsentinel.ai", "Platform Admin"),
    ("secadmin@netsentinel.ai", "Security Admin"),
    ("engineer@netsentinel.ai", "Network Engineer"),
    ("analyst@netsentinel.ai", "Security Analyst"),
    ("ciso@netsentinel.ai", "CISO"),
    ("auditor@netsentinel.ai", "Auditor"),
]:
    cells = table.add_row().cells
    cells[0].text, cells[1].text = email, role
body(doc, "See docs/demo.md for a two-minute walkthrough script.")
doc.add_page_break()

# ---- 10. Testing & CI ----
h1(doc, "10. Testing & CI")
mono_block(doc,
    "cd backend\n"
    "uv run pytest             # full suite (265 tests)\n"
    "uv run ruff check .       # lint\n"
    "uv run ruff format --check .\n"
    "uv run mypy app           # types\n\n"
    "cd ../frontend\n"
    "npm run lint              # oxlint\n"
    "npx tsc --noEmit          # types\n"
    "npm run build")
body(doc, "CI (.github/workflows/ci.yml) runs all of the above on every push and pull request.")
doc.add_page_break()

# ---- 11. Roadmap ----
h1(doc, "11. Roadmap (SP2-SP6)")
body(doc, "Every later phase extends this pipeline at a named seam — it does not replace it.")
bullets(doc, [
    "SP2 — additional vendors (parser + detector only)",
    "SP3 — additional compliance frameworks (rule pack data only)",
    "SP4 — adaptive AI training loop: approved unknown-construct mappings feed back into future parsing; AI evaluation/accuracy tracking; broader AI interpretation of passing/failing rule results in plain language",
    "SP5 — configuration drift detection, what-if hardening, risk graph visualization",
    "SP6 — tamper-evident evidence ledger; SSO/SAML/LDAP; Postgres/Redis/MinIO/Celery containerization (the three environment seams in Section 3.3 exist to make this a configuration change, not a rewrite)",
])

doc.save(OUT)
print(f"wrote {OUT}")
